"""Server-side XLSX/PDF generation for report exports. Deliberately generic
— every report type reduces to the same shape (a title, a list of column
headers, and a list of row tuples) before it reaches here, so export logic
lives in exactly one place instead of being duplicated per report."""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

# Column keys that hold money/decimal figures — right-aligned in the PDF,
# not specially formatted in Excel (Excel's own number formatting via cell
# type is enough; a raw Decimal serializes as a plain number in openpyxl).
NUMERIC_HINT_SUFFIXES = ("Amount", "Total", "Price", "Value", "Cost", "Revenue", "Profit", "Balance", "Qty", "%")


def _is_numeric_column(header: str) -> bool:
    return any(header.endswith(s) for s in NUMERIC_HINT_SUFFIXES)


def build_xlsx(title: str, columns: list[str], rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])

    for idx, header in enumerate(columns, start=1):
        col_letter = get_column_letter(idx)
        max_len = max([len(header)] + [len(str(r[idx - 1])) for r in rows]) if rows else len(header)
        ws.column_dimensions[col_letter].width = min(40, max(10, max_len + 2))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(title: str, columns: list[str], rows: list[list]) -> bytes:
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        leftMargin=12 * mm,
        rightMargin=12 * mm,
    )
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 8)]

    table_data = [columns] + [
        [str(v) if not isinstance(v, Decimal) else f"{v:,.2f}" for v in row] for row in rows
    ]
    table = Table(table_data, repeatRows=1)
    numeric_cols = {i for i, c in enumerate(columns) if _is_numeric_column(c)}
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    for col_idx in numeric_cols:
        style.append(("ALIGN", (col_idx, 0), (col_idx, -1), "RIGHT"))
    table.setStyle(TableStyle(style))
    elements.append(table)

    doc.build(elements)
    return buf.getvalue()
