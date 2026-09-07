"""Server-side XLSX/PDF generation for report exports. Deliberately generic
— every report type reduces to the same shape (a title, a list of column
headers, and a list of row tuples) before it reaches here, so export logic
lives in exactly one place instead of being duplicated per report."""

from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER

# Column keys that hold money/decimal figures — right-aligned in the PDF,
# not specially formatted in Excel (Excel's own number formatting via cell
# type is enough; a raw Decimal serializes as a plain number in openpyxl).
NUMERIC_HINT_SUFFIXES = ("Amount", "Total", "Price", "Value", "Cost", "Revenue", "Profit", "Balance", "Qty", "%")


def _is_numeric_column(header: str) -> bool:
    return any(header.endswith(s) for s in NUMERIC_HINT_SUFFIXES)


def build_xlsx(
    title: str,
    columns: list[str],
    rows: list[list],
    business_lines: list[str] | None = None,
) -> bytes:
    """business_lines (business name, PAN, address, contact — one string per
    row) prints as a small header block above the column headers, same
    identity block every report/statement export carries so a printed sheet
    is self-identifying without the on-screen app around it."""
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Report"

    header_row = 1
    if business_lines:
        for line in business_lines:
            ws.append([line])
            ws.cell(row=header_row, column=1).font = Font(bold=(header_row == 1), size=12 if header_row == 1 else 10)
            header_row += 1
        ws.append([])
        header_row += 1

    ws.append(columns)
    for cell in ws[header_row]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])

    for idx, header in enumerate(columns, start=1):
        col_letter = get_column_letter(idx)
        max_len = max([len(header)] + [len(str(r[idx - 1])) for r in rows]) if rows else len(header)
        ws.column_dimensions[col_letter].width = min(40, max(10, max_len + 2))

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(
    title: str,
    columns: list[str],
    rows: list[list],
    business_lines: list[str] | None = None,
    wide: bool = False,
) -> bytes:
    """A4 by default (portrait) — matches every other printed document in
    the app (see print.$invoiceId.tsx's A4 mode). Pass wide=True only for
    reports whose column count genuinely doesn't fit A4 portrait; that still
    prints on A4 paper, just landscape orientation, never a non-A4 size."""
    buf = BytesIO()
    pagesize = landscape(A4) if wide else A4
    margin = 10 * mm
    doc = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        topMargin=margin,
        bottomMargin=margin,
        leftMargin=margin,
        rightMargin=margin,
    )
    styles = getSampleStyleSheet()
    business_style = ParagraphStyle(
        "Business", parent=styles["Normal"], alignment=TA_CENTER, fontSize=9, leading=12
    )
    business_name_style = ParagraphStyle(
        "BusinessName", parent=styles["Heading2"], alignment=TA_CENTER, spaceAfter=2
    )
    title_style = ParagraphStyle(
        "ReportTitle", parent=styles["Heading3"], alignment=TA_CENTER, spaceBefore=4, spaceAfter=10
    )

    elements = []
    if business_lines:
        elements.append(Paragraph(business_lines[0], business_name_style))
        for line in business_lines[1:]:
            elements.append(Paragraph(line, business_style))
    elements.append(Paragraph(title, title_style))

    # A wide report (many columns, e.g. the 30-field Annex-5 Standard View)
    # needs a smaller font or every column collides with its neighbor —
    # plain-string table cells never wrap, they just overflow into the next
    # cell (reportlab's Table only wraps Paragraph flowables, not raw
    # strings). Scale font size down as column count grows, and wrap BOTH
    # header and data cells via Paragraph so long values break onto a
    # second line inside their own cell instead of overflowing into the
    # next one.
    font_size = 9.5 if len(columns) <= 10 else (8 if len(columns) <= 16 else (6.5 if len(columns) <= 22 else 5.5))
    header_style = ParagraphStyle(
        "TableHeader", parent=styles["Normal"], fontSize=font_size, leading=font_size + 2,
        textColor=colors.white, fontName="Helvetica-Bold", alignment=TA_CENTER,
    )
    cell_style = ParagraphStyle(
        "TableCell", parent=styles["Normal"], fontSize=font_size, leading=font_size + 2,
    )
    cell_style_right = ParagraphStyle(
        "TableCellRight", parent=cell_style, alignment=2,  # TA_RIGHT
    )
    numeric_cols = {i for i, c in enumerate(columns) if _is_numeric_column(c)}
    header_row = [Paragraph(str(h), header_style) for h in columns]
    table_data = [header_row] + [
        [
            Paragraph(
                str(v) if not isinstance(v, Decimal) else f"{v:,.2f}",
                cell_style_right if i in numeric_cols else cell_style,
            )
            for i, v in enumerate(row)
        ]
        for row in rows
    ]

    # Column widths default to content-width in reportlab, which routinely
    # leaves the table narrower than the page — the exact "empty margin on
    # the right" the printed output showed. Instead, weight each column by
    # its longest cell (header or data) and scale those weights to fill the
    # full printable width, so the table always spans edge-to-edge. Weights
    # come from the RAW values, not table_data — table_data now holds
    # Paragraph objects (needed so a long cell wraps instead of overflowing
    # into its neighbor), and str(Paragraph(...)) is not its rendered text.
    available_width = pagesize[0] - 2 * margin
    raw_str_rows = [
        [str(v) if not isinstance(v, Decimal) else f"{v:,.2f}" for v in row] for row in rows
    ]
    weights = [
        max([len(str(h))] + [len(r[i]) for r in raw_str_rows])
        for i, h in enumerate(columns)
    ]
    total_weight = sum(weights) or 1
    min_col_width = (9 if len(columns) > 22 else 10 if len(columns) > 16 else 12 if len(columns) > 10 else 15) * mm
    col_widths = [max(min_col_width, available_width * w / total_weight) for w in weights]
    # Rescale down if the minimum-width floor pushed the total over budget.
    scale = available_width / sum(col_widths)
    if scale < 1:
        col_widths = [w * scale for w in col_widths]

    table = Table(table_data, colWidths=col_widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 1), (-1, -1), font_size),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4 if len(columns) > 16 else 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4 if len(columns) > 16 else 5),
        ("LEFTPADDING", (0, 0), (-1, -1), 3 if len(columns) > 16 else 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 3 if len(columns) > 16 else 6),
    ]
    table.setStyle(TableStyle(style))
    elements.append(table)

    doc.build(elements)
    return buf.getvalue()
